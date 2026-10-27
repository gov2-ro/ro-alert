"""Build small synthetic RO-ALERT workbooks for tests (no personal messages, no real alerts).

Used by tests/fixtures/make_fixtures.py (also writes .xls, needs the development-only
xlwt package) and by the tests (.xlsx only, needs openpyxl from requirements.txt).
A workbook spec is {sheet name: list of rows}; a row is a list of cell texts.
"""

DETAIL_HEADER = ["", "Nume si data alerta", "Tip inregistrare", "Initiator alerta", "Risc principal",
                 "Risc asociat", "Data creare", "Data trimitere", "Status trimitere", "Tip mesaj",
                 "Lista poligoane", "Continut mesaj"]
REPORT_HEADER = ["Nr. crt.", "Nume si data alerta", "Tip inregistrare", "Initiator alerta",
                 "Risc principal", "Risc asociat", "Data trimitere", "Status trimitere",
                 "Lista poligoane", "Continut mesaj"]


def summary_sheet(title, items, total):
    """Rows of a summary sheet: items are (risc principal, risc asociat, total)."""
    rows = [["", "", "MINISTERUL TEST"], [], ["", "", "", title], [],
            ["", "", "Risc principal", "Risc asociat", "Total"]]
    last = None
    for principal, asociat, n in items:
        rows.append(["", "", principal if principal != last else "", asociat, str(n)])
        last = principal
    rows.append(["", "", "Grand Total", "", str(total)])
    return rows


def detail_rows(month):
    """Four delivery rows: alert A (SENT + ERROR retry), alert B, a CANCEL of B. Month 'YYYY/MM'."""
    def row(title, ch, init, asoc, created, sent, status, kind, poly, msg):
        return ["", title, "trimitere alerta pe canalul " + ch, init, "Altele", asoc,
                created, sent, status, kind, poly, msg]
    t = month
    return [
        row("ISU TEST - alerta A", "CB", "ISUAB", "Animale periculoase", t + "/01 10:00:00", t + "/01 10:00:05",
            "SENT", "ALERT", "1-ALBA IULIA", "Mesaj de test A. Test message A."),
        row("ISU TEST - alerta A", "ST", "ISUAB", "Animale periculoase", t + "/01 10:00:00", t + "/01 10:00:09",
            "ERROR", "ALERT", "1-ALBA IULIA", "Mesaj de test A. Test message A."),
        row("IGSU - alerta B", "CB", "IGSU", "Alte riscuri", t + "/02 11:00:00", t + "/02 11:00:03",
            "SENT", "ALERT", "", "Mesaj de test B. Test message B."),
        row("IGSU - alerta B", "CB", "IGSU", "Alte riscuri", t + "/03 12:00:00", t + "/03 12:00:04",
            "SENT", "CANCEL", "", "Anulare B. Cancel B."),
    ]


def report_rows(month, n=3):
    out = []
    for i in range(n):
        out.append([str(i + 1), "ISU TEST - alerta %d" % (i + 1), "trimitere alerta pe canalul CB",
                    "ISUAB", "Altele", "Alte riscuri", "%s/%02d 09:00:00" % (month, i + 1), "SENT",
                    "", "Mesaj de test %d. Test message %d." % (i + 1, i + 1)])
    return out


def detail_book(month="2030/03", header=DETAIL_HEADER, rows=None):
    rows = detail_rows(month) if rows is None else rows
    return {"Message_Details": [[], header] + rows}


def report_book(title, month="2030/04", n=3, total=None, header=REPORT_HEADER, rows=None):
    rows = report_rows(month, n) if rows is None else rows
    return {"Mesaje": summary_sheet(title, [("Altele", "Alte riscuri", len(rows))],
                                    len(rows) if total is None else total),
            "Raport": [[], header] + rows}


def summary_book(title, items, total):
    return {"Foaie1": summary_sheet(title, items, total), "Foaie2": [], "Foaie3": []}


def reorder(header, rows, order):
    """Return (header, rows) with the columns in `order` (a list of indexes)."""
    return [header[i] for i in order], [[r[i] for i in order] for r in rows]


def write_xlsx(path, book):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in book.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)


def write_xls(path, book):
    import xlwt  # development only; the committed fixtures are the test input
    wb = xlwt.Workbook()
    for name, rows in book.items():
        ws = wb.add_sheet(name)
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                ws.write(i, j, v)
    wb.save(path)
